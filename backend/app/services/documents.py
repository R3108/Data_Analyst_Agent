"""Stakeholder deliverables: real PDF documents and editable PowerPoint decks.

Charts are rebuilt as native objects in each format — vector shapes in the PDF and
genuine PowerPoint charts a colleague can restyle or re-point at new data — using
the numeric digest the sandbox already captures for every figure. No headless
browser, no image rasterisation, no extra binaries.
"""

from __future__ import annotations

import io
import logging
from typing import Any

from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.charts.linecharts import HorizontalLineChart
from reportlab.graphics.charts.piecharts import Pie
from reportlab.graphics.shapes import Drawing
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    KeepTogether,
    ListFlowable,
    ListItem,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

logger = logging.getLogger(__name__)

PALETTE_HEX = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300", "#4a3aa7", "#e34948"]
PALETTE = [colors.HexColor(c) for c in PALETTE_HEX]
INK = colors.HexColor("#0b0b0b")
INK_2 = colors.HexColor("#52514e")
INK_3 = colors.HexColor("#898781")
LINE = colors.HexColor("#dcdbd4")
MUTED = colors.HexColor("#f1f0ec")
GOOD = colors.HexColor("#006300")
BAD = colors.HexColor("#b83232")
WARN = colors.HexColor("#8a5a00")

MAX_TABLE_ROWS = 14
MAX_CHART_CATEGORIES = 12


# ============================================================================== shared helpers


def format_kpi(kpi: dict[str, Any]) -> str:
    value, fmt = kpi.get("value"), kpi.get("format")
    if value is None:
        return "—"
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (int, float)):
        number = float(value)
        if fmt == "percent":
            return f"{number * 100:,.1f}%"
        if fmt == "currency":
            return f"${number:,.0f}" if abs(number) >= 1000 else f"${number:,.2f}"
        if fmt == "integer" or float(number).is_integer():
            return f"{number:,.0f}"
        return f"{number:,.2f}"
    return str(value)


def format_delta(kpi: dict[str, Any]) -> str:
    delta = kpi.get("delta")
    if not isinstance(delta, (int, float)) or isinstance(delta, bool):
        return ""
    label = kpi.get("delta_label") or ""
    return f"{delta * 100:+.1f}% {label}".strip()


def _strip_markdown(text: str) -> str:
    """Flatten inline markdown to the light HTML subset reportlab understands."""
    import re

    out = str(text or "")
    out = re.sub(r"`([^`]+)`", r"\1", out)
    out = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", out)
    out = re.sub(r"(?<!\*)\*([^*\n]+)\*(?!\*)", r"<i>\1</i>", out)
    out = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", out)
    out = out.replace("&", "&amp;").replace("&amp;lt;", "&lt;")
    out = out.replace("<b>", "\x00b\x01").replace("</b>", "\x00/b\x01")
    out = out.replace("<i>", "\x00i\x01").replace("</i>", "\x00/i\x01")
    out = out.replace("<", "&lt;").replace(">", "&gt;")
    out = out.replace("\x00b\x01", "<b>").replace("\x00/b\x01", "</b>")
    out = out.replace("\x00i\x01", "<i>").replace("\x00/i\x01", "</i>")
    return out.strip()


def _paragraph_blocks(markdown: str) -> list[tuple[str, str]]:
    """Split an answer into ("p" | "li", text) blocks."""
    blocks: list[tuple[str, str]] = []
    for raw in str(markdown or "").split("\n"):
        line = raw.strip()
        if not line:
            continue
        if line.startswith(("- ", "* ", "• ")):
            blocks.append(("li", line[2:].strip()))
        elif line[:2].rstrip(".").isdigit() and line[1:3] in (". ", ") "):
            blocks.append(("li", line[3:].strip()))
        elif line.startswith("#"):
            blocks.append(("p", line.lstrip("# ").strip()))
        else:
            blocks.append(("p", line))
    return blocks


def _chart_series(chart: dict[str, Any]) -> dict[str, Any] | None:
    """Reduce a figure digest to categories plus one or more numeric series."""
    digest = chart.get("digest") or {}
    traces = [t for t in digest.get("traces") or [] if t]
    if not traces:
        return None

    kind = "bar"
    categories: list[str] = []
    series: list[tuple[str, list[float]]] = []

    for trace in traces[:4]:
        trace_type = str(trace.get("type") or "")
        if trace_type in ("scatter", "scattergl", "line"):
            kind = "line"
        elif trace_type == "pie":
            kind = "pie"

        if trace_type == "pie":
            labels = trace.get("labels") or []
            values = _numeric(trace.get("values") or [])
            if labels and values:
                categories = [str(v) for v in labels][:MAX_CHART_CATEGORIES]
                series = [(str(trace.get("name") or "Share"), values[:MAX_CHART_CATEGORIES])]
            continue

        ys = _numeric(trace.get("y") or [])
        xs = trace.get("x") or []
        if not ys:
            continue
        if not categories and xs:
            categories = [_short_label(x) for x in xs][:MAX_CHART_CATEGORIES]
        series.append((str(trace.get("name") or f"Series {len(series) + 1}"), ys[:MAX_CHART_CATEGORIES]))

    if not series:
        return None
    if not categories:
        categories = [str(i + 1) for i in range(len(series[0][1]))]
    width = min(len(categories), max(len(values) for _, values in series))
    if width < 1:
        return None
    return {
        "kind": kind,
        "categories": categories[:width],
        "series": [(name, values[:width]) for name, values in series if values],
        "axes": digest.get("axes") or {},
    }


def _numeric(values: list[Any]) -> list[float]:
    """Coerce a digest series to floats; non-numeric points become gaps at zero."""
    out: list[float] = []
    for value in values:
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            out.append(0.0)
        else:
            out.append(float(value))
    return out


def _short_label(value: Any) -> str:
    text = str(value)
    if len(text) >= 10 and text[4] == "-" and text[7] == "-":
        return text[:7] if text.endswith(("-01", "-01T00:00:00")) else text[:10]
    return text if len(text) <= 14 else text[:13] + "…"


# ============================================================================== PDF


def _styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "eyebrow": ParagraphStyle("eyebrow", parent=base["Normal"], fontName="Helvetica-Bold",
                                  fontSize=8, leading=11, textColor=INK_3, spaceAfter=4),
        "title": ParagraphStyle("title", parent=base["Title"], fontName="Helvetica-Bold",
                                fontSize=24, leading=28, textColor=INK, alignment=TA_LEFT,
                                spaceAfter=6),
        "meta": ParagraphStyle("meta", parent=base["Normal"], fontName="Helvetica", fontSize=9,
                               leading=13, textColor=INK_3, spaceAfter=14),
        "question": ParagraphStyle("question", parent=base["Heading2"], fontName="Helvetica-Bold",
                                   fontSize=14, leading=18, textColor=INK, spaceBefore=16,
                                   spaceAfter=6),
        "headline": ParagraphStyle("headline", parent=base["Normal"], fontName="Helvetica-Bold",
                                   fontSize=11.5, leading=16, textColor=INK, spaceAfter=8),
        "body": ParagraphStyle("body", parent=base["Normal"], fontName="Helvetica", fontSize=9.8,
                               leading=14.5, textColor=INK, spaceAfter=6),
        "bullet": ParagraphStyle("bullet", parent=base["Normal"], fontName="Helvetica", fontSize=9.8,
                                 leading=14, textColor=INK),
        "section": ParagraphStyle("section", parent=base["Normal"], fontName="Helvetica-Bold",
                                  fontSize=8.5, leading=11, textColor=INK_3, spaceBefore=10,
                                  spaceAfter=5),
        "caption": ParagraphStyle("caption", parent=base["Normal"], fontName="Helvetica-Oblique",
                                  fontSize=8.5, leading=11.5, textColor=INK_3, spaceAfter=10),
        "cell": ParagraphStyle("cell", parent=base["Normal"], fontName="Helvetica", fontSize=8,
                               leading=10.5, textColor=INK),
        "cellhead": ParagraphStyle("cellhead", parent=base["Normal"], fontName="Helvetica-Bold",
                                   fontSize=8, leading=10.5, textColor=INK_2),
    }


def session_to_pdf(session: dict[str, Any], messages: list[dict[str, Any]]) -> bytes:
    """A print-ready report of one analysis conversation."""
    styles = _styles()
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4, title=session["title"], author="Numera",
        leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm,
    )
    story: list[Any] = [
        Paragraph("ANALYSIS REPORT", styles["eyebrow"]),
        Paragraph(_strip_markdown(session["title"]), styles["title"]),
        Paragraph(
            f"{session.get('dataset_name', '')} · generated by Numera · every figure computed by "
            "executed code",
            styles["meta"],
        ),
    ]

    for message in messages:
        if message["role"] == "user":
            story.append(Paragraph(_strip_markdown(message["content"]), styles["question"]))
            continue
        story += _answer_flowables(message, styles)

    story.append(Spacer(1, 10))
    story.append(Paragraph(
        "Numera runs model-written Python in an isolated sandbox and verifies the write-up "
        "against the computed output before export.", styles["caption"]))
    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return buffer.getvalue()


def board_to_pdf(board: dict[str, Any]) -> bytes:
    """A print-ready dashboard of pinned items."""
    styles = _styles()
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4, title=board["title"], author="Numera",
        leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm,
    )
    story: list[Any] = [
        Paragraph("DASHBOARD", styles["eyebrow"]),
        Paragraph(_strip_markdown(board["title"]), styles["title"]),
    ]
    if board.get("description"):
        story.append(Paragraph(_strip_markdown(board["description"]), styles["meta"]))

    kpis = [item for item in board.get("items") or [] if item["kind"] == "kpi"]
    if kpis:
        story.append(Paragraph("KEY METRICS", styles["section"]))
        story.append(_kpi_table([{**item["content"], "label": item["title"]} for item in kpis], styles))
        story.append(Spacer(1, 8))

    for item in board.get("items") or []:
        if item["kind"] == "kpi":
            continue
        story += _board_item_flowables(item, styles)

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return buffer.getvalue()


def _footer(canvas: Any, doc: Any) -> None:
    canvas.saveState()
    canvas.setFont("Helvetica", 7.5)
    canvas.setFillColor(INK_3)
    canvas.drawString(18 * mm, 10 * mm, "Numera")
    canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Page {doc.page}")
    canvas.setStrokeColor(LINE)
    canvas.setLineWidth(0.5)
    canvas.line(18 * mm, 13.5 * mm, A4[0] - 18 * mm, 13.5 * mm)
    canvas.restoreState()


def _answer_flowables(message: dict[str, Any], styles: dict[str, ParagraphStyle]) -> list[Any]:
    payload = message.get("payload") or {}
    report = payload.get("report") or {}
    execution = payload.get("execution") or {}
    out: list[Any] = []

    if report.get("headline"):
        out.append(Paragraph(_strip_markdown(report["headline"]), styles["headline"]))

    body = report.get("answer_markdown") or message.get("content") or ""
    out += _markdown_flowables(body, styles)

    verification = payload.get("verification")
    if verification and verification.get("score") is not None:
        out.append(_verification_flowable(verification, styles))

    if execution.get("kpis"):
        out.append(Paragraph("KEY METRICS", styles["section"]))
        out.append(_kpi_table(execution["kpis"], styles))

    for chart in execution.get("charts") or []:
        drawing = _pdf_chart(chart)
        if drawing is None:
            continue
        block: list[Any] = [Paragraph(_strip_markdown(chart["title"]), styles["section"]), drawing]
        if chart.get("caption"):
            block.append(Paragraph(_strip_markdown(chart["caption"]), styles["caption"]))
        out.append(KeepTogether(block))

    for table in execution.get("tables") or []:
        out.append(KeepTogether([
            Paragraph(_strip_markdown(table["title"]), styles["section"]),
            _data_table(table, styles),
        ]))

    if report.get("insights"):
        out.append(Paragraph("INSIGHTS", styles["section"]))
        out.append(ListFlowable(
            [ListItem(Paragraph(f"<b>{_strip_markdown(i['title'])}</b> — {_strip_markdown(i['detail'])}",
                                styles["bullet"]), leftIndent=12) for i in report["insights"]],
            bulletType="bullet", bulletFontSize=6, leftIndent=10, spaceAfter=6,
        ))
    if report.get("recommendations"):
        out.append(Paragraph("RECOMMENDED ACTIONS", styles["section"]))
        out.append(ListFlowable(
            [ListItem(Paragraph(_strip_markdown(r), styles["bullet"]), leftIndent=12)
             for r in report["recommendations"]],
            bulletType="1", leftIndent=10, spaceAfter=6,
        ))
    if report.get("caveats"):
        out.append(Paragraph("CAVEATS", styles["section"]))
        out.append(ListFlowable(
            [ListItem(Paragraph(_strip_markdown(c), styles["bullet"]), leftIndent=12)
             for c in report["caveats"]],
            bulletType="bullet", bulletFontSize=6, leftIndent=10, spaceAfter=6,
        ))
    return out


def _markdown_flowables(markdown: str, styles: dict[str, ParagraphStyle]) -> list[Any]:
    out: list[Any] = []
    bullets: list[str] = []

    def flush() -> None:
        if bullets:
            out.append(ListFlowable(
                [ListItem(Paragraph(text, styles["bullet"]), leftIndent=12) for text in bullets],
                bulletType="bullet", bulletFontSize=6, leftIndent=10, spaceAfter=6,
            ))
            bullets.clear()

    for kind, text in _paragraph_blocks(markdown):
        rendered = _strip_markdown(text)
        if not rendered:
            continue
        if kind == "li":
            bullets.append(rendered)
        else:
            flush()
            out.append(Paragraph(rendered, styles["body"]))
    flush()
    return out


def _verification_flowable(verification: dict[str, Any], styles: dict[str, ParagraphStyle]) -> Any:
    confidence = verification.get("confidence", "unverified")
    tone = {"high": GOOD, "medium": WARN, "low": BAD}.get(confidence, INK_3)
    lines = [
        f"<b>Verification: {confidence} confidence ({verification['score']}/100)</b> · "
        f"{verification.get('checks', 0)} automated checks"
    ]
    for finding in verification.get("findings") or []:
        lines.append(f"• {_strip_markdown(finding['title'])} ({finding['severity']}) — "
                     f"{_strip_markdown(finding['detail'])}")
    cell = [[Paragraph("<br/>".join(lines), styles["cell"])]]
    table = Table(cell, colWidths=[None])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), MUTED),
        ("LINEBEFORE", (0, 0), (0, -1), 2, tone),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    return table


def _kpi_table(kpis: list[dict[str, Any]], styles: dict[str, ParagraphStyle]) -> Table:
    columns = min(len(kpis), 4) or 1
    cells: list[list[Any]] = []
    for start in range(0, len(kpis), columns):
        chunk = kpis[start : start + columns]
        labels = [Paragraph(_strip_markdown(k.get("label", "")), styles["cellhead"]) for k in chunk]
        values = [Paragraph(f"<b>{format_kpi(k)}</b>", styles["body"]) for k in chunk]
        deltas = [Paragraph(format_delta(k), styles["cell"]) for k in chunk]
        padding = [""] * (columns - len(chunk))
        cells += [labels + padding, values + padding, deltas + padding]

    table = Table(cells, colWidths=[(A4[0] - 36 * mm) / columns] * columns)
    style = [
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]
    for group in range(len(cells) // 3):
        top, bottom = group * 3, group * 3 + 2
        style += [
            ("BACKGROUND", (0, top), (-1, bottom), MUTED),
            ("TOPPADDING", (0, top), (-1, top), 7),
            ("BOTTOMPADDING", (0, bottom), (-1, bottom), 7),
        ]
    table.setStyle(TableStyle(style))
    return table


def _data_table(table_output: dict[str, Any], styles: dict[str, ParagraphStyle]) -> Table:
    headers = [c["name"] for c in table_output.get("columns") or []]
    rows = (table_output.get("rows") or [])[:MAX_TABLE_ROWS]
    data = [[Paragraph(_strip_markdown(h), styles["cellhead"]) for h in headers]]
    for row in rows:
        data.append([Paragraph(_cell_text(v), styles["cell"]) for v in row])

    available = A4[0] - 36 * mm
    table = Table(data, colWidths=[available / max(len(headers), 1)] * len(headers), repeatRows=1)
    table.setStyle(TableStyle([
        ("LINEBELOW", (0, 0), (-1, 0), 0.75, LINE),
        ("LINEBELOW", (0, 1), (-1, -2), 0.25, LINE),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#faf9f6")]),
    ]))
    return table


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        return f"{value:,.2f}" if not float(value).is_integer() else f"{value:,.0f}"
    if isinstance(value, int):
        return f"{value:,}"
    text = str(value)
    if len(text) >= 19 and text[4] == "-" and "T" in text:
        return text[:10]
    return _strip_markdown(text[:60])


def _pdf_chart(chart: dict[str, Any]) -> Drawing | None:
    """Rebuild a figure as reportlab vector graphics from its numeric digest."""
    try:
        spec = _chart_series(chart)
        if not spec:
            return None
        width, height = A4[0] - 36 * mm, 62 * mm
        drawing = Drawing(width, height)
        categories, series = spec["categories"], spec["series"]

        if spec["kind"] == "pie":
            pie = Pie()
            pie.x, pie.y = 12, 8
            pie.width = pie.height = height - 18
            pie.data = [abs(v) for v in series[0][1]]
            pie.labels = [f"{c}" for c in categories]
            pie.slices.strokeWidth = 0.5
            pie.slices.strokeColor = colors.white
            pie.sideLabels = True
            for index in range(len(pie.data)):
                pie.slices[index].fillColor = PALETTE[index % len(PALETTE)]
            drawing.add(pie)
            return drawing

        chart_obj: Any
        if spec["kind"] == "line":
            chart_obj = HorizontalLineChart()
            chart_obj.lines.strokeWidth = 1.6
            for index in range(len(series)):
                chart_obj.lines[index].strokeColor = PALETTE[index % len(PALETTE)]
        else:
            chart_obj = VerticalBarChart()
            chart_obj.barSpacing = 1
            chart_obj.groupSpacing = 6
            for index in range(len(series)):
                chart_obj.bars[index].fillColor = PALETTE[index % len(PALETTE)]
                chart_obj.bars[index].strokeColor = None

        chart_obj.x, chart_obj.y = 34, 26
        chart_obj.width = width - 50
        chart_obj.height = height - 40
        chart_obj.data = [values for _, values in series]
        chart_obj.categoryAxis.categoryNames = categories
        chart_obj.categoryAxis.labels.fontName = "Helvetica"
        chart_obj.categoryAxis.labels.fontSize = 6.5
        chart_obj.categoryAxis.labels.angle = 30 if max(len(c) for c in categories) > 6 else 0
        chart_obj.categoryAxis.labels.dy = -6
        chart_obj.categoryAxis.labels.boxAnchor = "ne" if chart_obj.categoryAxis.labels.angle else "n"
        chart_obj.categoryAxis.strokeColor = LINE
        chart_obj.valueAxis.labels.fontName = "Helvetica"
        chart_obj.valueAxis.labels.fontSize = 6.5
        chart_obj.valueAxis.strokeColor = LINE
        chart_obj.valueAxis.gridStrokeColor = LINE
        chart_obj.valueAxis.gridStrokeWidth = 0.25
        chart_obj.valueAxis.visibleGrid = True

        flat = [v for _, values in series for v in values]
        low, high = min(flat + [0.0]), max(flat + [0.0])
        if high == low:
            high = low + 1
        chart_obj.valueAxis.valueMin = low if low < 0 else 0
        chart_obj.valueAxis.valueMax = high * 1.08
        drawing.add(chart_obj)
        return drawing
    except Exception:  # noqa: BLE001 — a chart must never break an export
        logger.debug("Could not render chart '%s' into the PDF", chart.get("title"), exc_info=True)
        return None


def _board_item_flowables(item: dict[str, Any], styles: dict[str, ParagraphStyle]) -> list[Any]:
    content = item.get("content") or {}
    if item["kind"] == "chart":
        drawing = _pdf_chart(content)
        if drawing is None:
            return []
        block = [Paragraph(_strip_markdown(item["title"]), styles["section"]), drawing]
        if content.get("caption"):
            block.append(Paragraph(_strip_markdown(content["caption"]), styles["caption"]))
        return [KeepTogether(block)]
    if item["kind"] == "table":
        return [KeepTogether([
            Paragraph(_strip_markdown(item["title"]), styles["section"]),
            _data_table(content, styles),
        ])]
    if item["kind"] == "insight":
        return [Paragraph(f"<b>{_strip_markdown(item['title'])}</b> — "
                          f"{_strip_markdown(content.get('detail', ''))}", styles["body"])]
    if item["kind"] == "note":
        return _markdown_flowables(content.get("text", ""), styles)
    return []


# ============================================================================== PowerPoint


def session_to_pptx(session: dict[str, Any], messages: list[dict[str, Any]]) -> bytes:
    """An editable deck: one section per question, native charts a colleague can restyle."""
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    _title_slide(presentation, session["title"],
                 f"{session.get('dataset_name', '')} · analysed with Numera")

    question = ""
    for message in messages:
        if message["role"] == "user":
            question = message["content"]
            continue
        payload = message.get("payload") or {}
        _answer_slides(presentation, question, payload, message.get("content", ""))

    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def board_to_pptx(board: dict[str, Any]) -> bytes:
    """An editable deck of the pinned dashboard items."""
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    _title_slide(presentation, board["title"], board.get("description") or "Dashboard from Numera")

    items = board.get("items") or []
    kpis = [{**i["content"], "label": i["title"]} for i in items if i["kind"] == "kpi"]
    if kpis:
        slide = _content_slide(presentation, "Key metrics")
        _kpi_row(slide, kpis)

    for item in items:
        content = item.get("content") or {}
        if item["kind"] == "chart":
            slide = _content_slide(presentation, item["title"])
            _native_chart(slide, content)
        elif item["kind"] == "table":
            slide = _content_slide(presentation, item["title"])
            _pptx_table(slide, content)
        elif item["kind"] in ("insight", "note"):
            slide = _content_slide(presentation, item["title"])
            text = content.get("detail") or content.get("text") or ""
            _bullets(slide, [b for _, b in _paragraph_blocks(text)][:8])

    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _pptx_colors() -> Any:
    from pptx.dml.color import RGBColor

    return RGBColor


def _title_slide(presentation: Any, title: str, subtitle: str) -> Any:
    from pptx.util import Inches, Pt

    RGBColor = _pptx_colors()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.9), Inches(2.4), Inches(11.5), Inches(2.4))
    frame = box.text_frame
    frame.word_wrap = True
    frame.text = title
    run = frame.paragraphs[0].runs[0]
    run.font.size = Pt(40)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x0B, 0x0B, 0x0B)

    paragraph = frame.add_paragraph()
    paragraph.text = subtitle
    paragraph.runs[0].font.size = Pt(16)
    paragraph.runs[0].font.color.rgb = RGBColor(0x89, 0x87, 0x81)
    return slide


def _content_slide(presentation: Any, title: str) -> Any:
    from pptx.util import Inches, Pt

    RGBColor = _pptx_colors()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12.1), Inches(0.9))
    frame = box.text_frame
    frame.word_wrap = True
    frame.text = str(title)[:160]
    run = frame.paragraphs[0].runs[0]
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x0B, 0x0B, 0x0B)
    return slide


def _answer_slides(presentation: Any, question: str, payload: dict[str, Any], fallback: str) -> None:
    report = payload.get("report") or {}
    execution = payload.get("execution") or {}

    slide = _content_slide(presentation, question or report.get("headline") or "Analysis")
    lines: list[str] = []
    if report.get("headline"):
        lines.append(report["headline"])
    body = report.get("answer_markdown") or fallback
    lines += [text for _, text in _paragraph_blocks(body)][:5]
    _bullets(slide, lines[:6])
    if execution.get("kpis"):
        _kpi_row(slide, execution["kpis"], top=5.3)

    verification = payload.get("verification")
    if verification and verification.get("score") is not None:
        _footnote(slide, f"Verification: {verification['confidence']} confidence "
                         f"({verification['score']}/100), {verification.get('checks', 0)} automated checks")

    for chart in execution.get("charts") or []:
        chart_slide = _content_slide(presentation, chart["title"])
        if _native_chart(chart_slide, chart) and chart.get("caption"):
            _footnote(chart_slide, chart["caption"])

    for table in execution.get("tables") or []:
        table_slide = _content_slide(presentation, table["title"])
        _pptx_table(table_slide, table)

    if report.get("insights") or report.get("recommendations"):
        summary = _content_slide(presentation, "Insights and next steps")
        bullets = [f"{i['title']} — {i['detail']}" for i in report.get("insights") or []]
        bullets += [f"Action: {r}" for r in report.get("recommendations") or []]
        _bullets(summary, bullets[:8])


def _bullets(slide: Any, lines: list[str], top: float = 1.5) -> None:
    from pptx.util import Inches, Pt

    RGBColor = _pptx_colors()
    if not lines:
        return
    box = slide.shapes.add_textbox(Inches(0.7), Inches(top), Inches(11.9), Inches(3.6))
    frame = box.text_frame
    frame.word_wrap = True
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = _plain(line)[:400]
        paragraph.space_after = Pt(8)
        for run in paragraph.runs:
            run.font.size = Pt(15 if index == 0 else 13)
            run.font.bold = index == 0
            run.font.color.rgb = RGBColor(0x0B, 0x0B, 0x0B) if index == 0 else RGBColor(0x52, 0x51, 0x4E)


def _footnote(slide: Any, text: str) -> None:
    from pptx.util import Inches, Pt

    RGBColor = _pptx_colors()
    box = slide.shapes.add_textbox(Inches(0.7), Inches(6.85), Inches(11.9), Inches(0.4))
    frame = box.text_frame
    frame.word_wrap = True
    frame.text = _plain(text)[:220]
    run = frame.paragraphs[0].runs[0]
    run.font.size = Pt(10)
    run.font.italic = True
    run.font.color.rgb = RGBColor(0x89, 0x87, 0x81)


def _kpi_row(slide: Any, kpis: list[dict[str, Any]], top: float = 2.0) -> None:
    from pptx.util import Inches, Pt

    RGBColor = _pptx_colors()
    shown = kpis[:5]
    if not shown:
        return
    width = 11.9 / len(shown)
    for index, kpi in enumerate(shown):
        box = slide.shapes.add_textbox(Inches(0.7 + index * width), Inches(top),
                                       Inches(width - 0.25), Inches(1.4))
        frame = box.text_frame
        frame.word_wrap = True
        frame.text = str(kpi.get("label", ""))[:60]
        label_run = frame.paragraphs[0].runs[0]
        label_run.font.size = Pt(11)
        label_run.font.color.rgb = RGBColor(0x52, 0x51, 0x4E)

        value = frame.add_paragraph()
        value.text = format_kpi(kpi)
        value.runs[0].font.size = Pt(26)
        value.runs[0].font.bold = True
        value.runs[0].font.color.rgb = RGBColor(0x0B, 0x0B, 0x0B)

        delta = format_delta(kpi)
        if delta:
            row = frame.add_paragraph()
            row.text = delta
            positive = not delta.startswith("-")
            good = positive == bool(kpi.get("higher_is_better", True))
            row.runs[0].font.size = Pt(11)
            row.runs[0].font.color.rgb = (
                RGBColor(0x00, 0x63, 0x00) if good else RGBColor(0xB8, 0x32, 0x32)
            )


def _native_chart(slide: Any, chart: dict[str, Any]) -> bool:
    """Add a real PowerPoint chart (editable, with its data embedded)."""
    try:
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
        from pptx.util import Inches, Pt

        spec = _chart_series(chart)
        if not spec:
            return False
        data = CategoryChartData()
        data.categories = spec["categories"]
        for name, values in spec["series"]:
            data.add_series(name, values)

        chart_type = {
            "line": XL_CHART_TYPE.LINE_MARKERS,
            "pie": XL_CHART_TYPE.PIE,
        }.get(spec["kind"], XL_CHART_TYPE.COLUMN_CLUSTERED)

        frame = slide.shapes.add_chart(
            chart_type, Inches(0.7), Inches(1.4), Inches(11.9), Inches(5.2), data
        )
        native = frame.chart
        native.has_title = False
        if len(spec["series"]) > 1 or spec["kind"] == "pie":
            native.has_legend = True
            native.legend.position = XL_LEGEND_POSITION.BOTTOM
            native.legend.include_in_layout = False
        else:
            native.has_legend = False
        native.font.size = Pt(11)
        return True
    except Exception:  # noqa: BLE001 — a chart must never break an export
        logger.debug("Could not render chart '%s' into the deck", chart.get("title"), exc_info=True)
        return False


def _pptx_table(slide: Any, table_output: dict[str, Any]) -> None:
    from pptx.util import Inches, Pt

    RGBColor = _pptx_colors()
    headers = [c["name"] for c in table_output.get("columns") or []]
    rows = (table_output.get("rows") or [])[:10]
    if not headers:
        return
    shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(0.7), Inches(1.5),
                                   Inches(11.9), Inches(0.4 + 0.32 * len(rows)))
    table = shape.table
    for column, name in enumerate(headers):
        cell = table.cell(0, column)
        cell.text = str(name)[:40]
        for paragraph in cell.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(11)
                run.font.bold = True
    for row_index, row in enumerate(rows, start=1):
        for column in range(len(headers)):
            cell = table.cell(row_index, column)
            cell.text = _cell_text(row[column]) if column < len(row) else ""
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)


def _plain(text: str) -> str:
    import re

    out = re.sub(r"`([^`]+)`", r"\1", str(text or ""))
    out = re.sub(r"\*\*([^*]+)\*\*", r"\1", out)
    out = re.sub(r"(?<!\*)\*([^*\n]+)\*(?!\*)", r"\1", out)
    out = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", out)
    return out.strip()
